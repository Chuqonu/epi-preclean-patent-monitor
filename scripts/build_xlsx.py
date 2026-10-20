#!/usr/bin/env python3
"""Build ledger/patents.xlsx from data/report.json (handoff 2026-10-20.4.2-r2).

Sheets:
  主表       — month EventRecords (core cards)
  同族事件   — family_events expansion
  关联待核   — pending_or_related + non-core statuses from months
  检索覆盖   — search_coverage
  变更记录   — change_history (top-level + per-record)

Usage:
  python3 scripts/build_xlsx.py
  python3 scripts/build_xlsx.py --input data/report.json --output ledger/patents.xlsx
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_IN = ROOT / "data" / "report.json"
DEFAULT_OUT = ROOT / "ledger" / "patents.xlsx"

MAIN_HEADERS = [
    "month",
    "id",
    "publication_id",
    "jurisdiction",
    "event_type",
    "event_date",
    "application_id",
    "priority_date",
    "family_id",
    "family_id_source",
    "title_zh",
    "title_original",
    "applicant_normalized",
    "applicant_original",
    "scope_status",
    "technology_tags",
    "short_preview",
    "technical_summary",
    "claims_or_embodiment_basis",
    "scope_reason",
    "family_completion_status",
    "source_urls",
    "source_retrieved_at",
    "source_sha256",
    "actual_read_spans",
    "reviewer",
    "reviewed_at",
    "figure_count",
    "import_tag",
    "discovery_source",
]

EVENT_HEADERS = [
    "month",
    "parent_publication_id",
    "family_id",
    "publication_id",
    "event_date",
    "event_type",
    "source_url",
    "application_id",
]

PENDING_HEADERS = MAIN_HEADERS[:]  # same shape

COVERAGE_HEADERS = [
    "jurisdiction",
    "branch",
    "window",
    "status",
    "notes",
]

CHANGE_HEADERS = [
    "scope",
    "record_id",
    "at",
    "by",
    "action",
    "detail",
]


def _get(rec: dict, *keys, default=None):
    for k in keys:
        if k in rec and rec[k] is not None:
            return rec[k]
    return default


def normalize_record(rec: dict, month: str | None = None) -> dict:
    """Map handoff fields; accept legacy aliases from reference report."""
    pub = _get(rec, "publication_id", "publication_number", default="")
    m = month or rec.get("month") or ""
    rid = rec.get("id") or (f"{m}-{pub}" if m and pub else pub)
    tags = _get(rec, "technology_tags", default=[]) or []
    if isinstance(tags, str):
        tags = [tags]
    urls = _get(rec, "source_urls", default=[]) or []
    spans = _get(rec, "actual_read_spans", default=[]) or []
    figs = _get(rec, "figures", "images", default=[]) or []
    return {
        "month": m,
        "id": rid,
        "publication_id": pub,
        "jurisdiction": _get(rec, "jurisdiction", default=""),
        "event_type": _get(rec, "event_type", default=""),
        "event_date": _get(rec, "event_date", "publication_date", default=""),
        "application_id": _get(rec, "application_id", "application_number", default=""),
        "priority_date": (
            rec.get("priority_date")
            if rec.get("priority_date") is not None
            else (rec.get("priority") or {}).get("date") if isinstance(rec.get("priority"), dict) else ""
        ),
        "family_id": _get(rec, "family_id", "family_key", default=""),
        "family_id_source": _get(rec, "family_id_source", default=""),
        "title_zh": _get(rec, "title_zh", "title", default=""),
        "title_original": _get(rec, "title_original", default=""),
        "applicant_normalized": _get(rec, "applicant_normalized", "applicant", default=""),
        "applicant_original": _get(rec, "applicant_original", default=""),
        "scope_status": _get(rec, "scope_status", default=""),
        "technology_tags": "; ".join(tags),
        "short_preview": _get(rec, "short_preview", "brief", default=""),
        "technical_summary": _get(rec, "technical_summary", "summary", default=""),
        "claims_or_embodiment_basis": _get(rec, "claims_or_embodiment_basis", "claims", default=""),
        "scope_reason": _get(rec, "scope_reason", "relevance", default=""),
        "family_completion_status": _get(rec, "family_completion_status", default=""),
        "source_urls": " | ".join(urls) if isinstance(urls, list) else str(urls or ""),
        "source_retrieved_at": _get(rec, "source_retrieved_at", default=""),
        "source_sha256": _get(rec, "source_sha256", default=""),
        "actual_read_spans": "; ".join(spans) if isinstance(spans, list) else str(spans or ""),
        "reviewer": _get(rec, "reviewer", default=""),
        "reviewed_at": _get(rec, "reviewed_at", default=""),
        "figure_count": len(figs),
        "import_tag": _get(rec, "import_tag", default=""),
        "discovery_source": _get(rec, "discovery_source", default=""),
        "_raw": rec,
    }


def family_events_of(rec: dict) -> list:
    return _get(rec, "family_events", "events", default=[]) or []


def write_sheet(ws, headers, rows):
    from openpyxl.styles import Font

    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    for col in ws.columns:
        max_len = 0
        letter = col[0].column_letter
        for cell in col[:200]:
            val = "" if cell.value is None else str(cell.value)
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[letter].width = min(48, max(10, max_len + 2))


def build(data: dict, out_path: Path) -> None:
    try:
        from openpyxl import Workbook
    except ImportError:
        print("openpyxl required: pip install openpyxl", file=sys.stderr)
        raise

    main_rows = []
    event_rows = []
    pending_rows = []
    change_rows = []

    for ch in data.get("change_history") or []:
        change_rows.append(
            {
                "scope": "report",
                "record_id": "",
                "at": ch.get("at", ""),
                "by": ch.get("by", ""),
                "action": ch.get("action", ""),
                "detail": ch.get("detail", ""),
            }
        )

    for month in data.get("months") or []:
        mlabel = month.get("month", "")
        for rec in month.get("families") or []:
            norm = normalize_record(rec, mlabel)
            main_rows.append(norm)
            status = (norm.get("scope_status") or "").lower()
            if status and status not in ("core", ""):
                pending_rows.append(norm)
            for ev in family_events_of(rec):
                event_rows.append(
                    {
                        "month": mlabel,
                        "parent_publication_id": norm["publication_id"],
                        "family_id": norm["family_id"],
                        "publication_id": _get(ev, "publication_id", "number", default=""),
                        "event_date": _get(ev, "event_date", "date", default=""),
                        "event_type": _get(ev, "event_type", "event", default=""),
                        "source_url": _get(ev, "source_url", default=""),
                        "application_id": _get(ev, "application_id", "application_number", default=""),
                    }
                )
            for ch in rec.get("change_history") or []:
                change_rows.append(
                    {
                        "scope": "record",
                        "record_id": norm["id"],
                        "at": ch.get("at", ""),
                        "by": ch.get("by", ""),
                        "action": ch.get("action", ""),
                        "detail": ch.get("detail", ""),
                    }
                )

    for rec in data.get("pending_or_related") or []:
        pending_rows.append(normalize_record(rec, rec.get("month")))

    coverage_rows = []
    for row in data.get("search_coverage") or []:
        coverage_rows.append(
            {
                "jurisdiction": row.get("jurisdiction", ""),
                "branch": row.get("branch", ""),
                "window": row.get("window", ""),
                "status": row.get("status", ""),
                "notes": row.get("notes", ""),
            }
        )

    # Always emit header-only sheets so the workbook shape is stable in scaffold.
    wb = Workbook()
    ws_main = wb.active
    ws_main.title = "主表"
    write_sheet(ws_main, MAIN_HEADERS, main_rows)

    ws_ev = wb.create_sheet("同族事件")
    write_sheet(ws_ev, EVENT_HEADERS, event_rows)

    ws_pend = wb.create_sheet("关联待核")
    write_sheet(ws_pend, PENDING_HEADERS, pending_rows)

    ws_cov = wb.create_sheet("检索覆盖")
    write_sheet(ws_cov, COVERAGE_HEADERS, coverage_rows)

    ws_ch = wb.create_sheet("变更记录")
    write_sheet(ws_ch, CHANGE_HEADERS, change_rows)

    # Meta sheet with honest counts
    ws_meta = wb.create_sheet("元数据", 0)
    counts = data.get("counts") or {}
    meta = [
        ("title", data.get("title", "")),
        ("period", data.get("period", "")),
        ("prepared_date", data.get("prepared_date") or ""),
        ("research_through", data.get("research_through") or ""),
        ("method_trace", data.get("method_trace") or ""),
        ("scope", data.get("scope", "")),
        ("excluded_scope", data.get("excluded_scope", "")),
        ("status_note", data.get("status_note", "")),
    ]
    for k, v in counts.items():
        meta.append((f"counts.{k}", v))
    meta.append(("note", "Empty months ≠ zero hits. 2025-09 paused. No zero-omission claim."))
    ws_meta.append(["key", "value"])
    from openpyxl.styles import Font

    for c in ws_meta[1]:
        c.font = Font(bold=True)
    for k, v in meta:
        ws_meta.append([k, v])
    ws_meta.column_dimensions["A"].width = 28
    ws_meta.column_dimensions["B"].width = 80

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"Wrote {out_path}")
    print(
        f"  主表={len(main_rows)} 同族事件={len(event_rows)} "
        f"关联待核={len(pending_rows)} 检索覆盖={len(coverage_rows)} 变更={len(change_rows)}"
    )


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--input", type=Path, default=DEFAULT_IN)
    p.add_argument("--output", type=Path, default=DEFAULT_OUT)
    args = p.parse_args(argv)
    if not args.input.is_file():
        print(f"Missing input: {args.input}", file=sys.stderr)
        return 1
    with args.input.open(encoding="utf-8") as f:
        data = json.load(f)
    build(data, args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
